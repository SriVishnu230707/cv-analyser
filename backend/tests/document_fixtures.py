"""Actual synthetic files for parser tests and browser QA."""
from io import BytesIO
import pymupdf
from docx import Document

RESUME = "Synthetic Candidate\nSummary\nSoftware developer building useful tools.\nSkills\nPython, SQL, Flask, C++, C#\nProjects\nBuilt REST APIs using Python and Flask.\nEducation\nCompleted an introductory programming course."


def pdf_bytes(pages=1, encrypted=False, columns=False):
    with pymupdf.open() as document:
        for _ in range(pages):
            page = document.new_page()
            if columns:
                page.insert_text((40, 70), "Skills\nPython, SQL, Flask", fontsize=13)
                page.insert_text((330, 70), "Experience\nBuilt REST APIs using Flask.\nQueried application data with SQL.", fontsize=13)
            else:
                page.insert_text((40, 70), RESUME, fontsize=13)
        if encrypted:
            return document.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="test-secret", owner_pw="test-owner")
        return document.tobytes()


def scanned_pdf(mixed=False):
    with pymupdf.open(stream=pdf_bytes(), filetype="pdf") as original:
        png = original[0].get_pixmap(dpi=180).tobytes("png")
    with pymupdf.open() as document:
        if mixed:
            with pymupdf.open(stream=pdf_bytes(), filetype="pdf") as original:
                document.insert_pdf(original)
        page = document.new_page()
        page.insert_image(page.rect, stream=png)
        return document.tobytes(deflate=True)


def docx_bytes():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Synthetic Candidate · fictional contact header"
    document.add_heading("Summary", level=1)
    document.add_paragraph("Software developer building useful tools.")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python, SQL, Flask, C++, C#")
    document.add_heading("Experience", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Backend project"
    table.cell(0, 1).text = "Built REST APIs using Python and Flask."
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Queried application data using SQL."
    nested = table.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested project detail retained."
    document.add_heading("Education", level=1)
    document.add_paragraph("Completed an introductory programming course.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
